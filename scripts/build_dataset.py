from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


COUNTIES = {
    "卢氏": "Lushi County",
    "鲁山": "Lushan County",
    "栾川": "Luanchuan County",
    "洛宁": "Luoning County",
    "南召": "Nanzhao County",
    "内乡": "Neixiang County",
    "汝阳": "Ruyang County",
    "嵩县": "Song County",
    "西峡": "Xixia County",
    "淅川": "Xichuan County",
    "宜阳": "Yiyang County",
}
SERVICES = {
    "供给": ("provisioning", "Provisioning"),
    "水源": ("water_retention", "Water retention"),
    "碳储": ("carbon_storage", "Carbon storage"),
    "土保": ("soil_retention", "Soil retention"),
    "旅游": ("recreation", "Recreation"),
}


def extract(source: Path) -> pd.DataFrame:
    worksheet = load_workbook(source, data_only=True, read_only=True)["Sheet1"]
    blocks = {2018: range(2, 13), 2023: range(19, 30)}
    records: list[dict[str, object]] = []

    for year, rows in blocks.items():
        header_row = 1 if year == 2018 else 17
        for row in rows:
            county_zh = worksheet.cell(row, 2).value
            if county_zh not in COUNTIES:
                raise ValueError(f"Unexpected county in row {row}: {county_zh!r}")
            for column in range(3, 8):
                service_zh = worksheet.cell(header_row, column).value
                if service_zh not in SERVICES:
                    raise ValueError(f"Unexpected service in column {column}: {service_zh!r}")
                service, service_label = SERVICES[service_zh]
                value_yuan = float(worksheet.cell(row, column).value)
                records.append(
                    {
                        "year": year,
                        "county": COUNTIES[county_zh],
                        "county_zh": county_zh,
                        "service": service,
                        "service_label": service_label,
                        "service_zh": service_zh,
                        "value_yuan": round(value_yuan, 2),
                        "value_cny_100m": round(value_yuan / 100_000_000, 8),
                    }
                )

    data = pd.DataFrame.from_records(records)
    if len(data) != 110:
        raise ValueError(f"Expected 110 records, found {len(data)}")
    return data


def main() -> None:
    parser = argparse.ArgumentParser(description="Migrate thesis results into a tidy CSV.")
    parser.add_argument("--source-workbook", required=True, type=Path)
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("data/processed/ecosystem_values_2018_2023.csv"),
    )
    args = parser.parse_args()
    data = extract(args.source_workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    data.to_csv(args.output, index=False, encoding="utf-8")
    print(f"Wrote {len(data)} records to {args.output}")


if __name__ == "__main__":
    main()
